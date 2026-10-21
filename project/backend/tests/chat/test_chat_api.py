import io
import json

import httpx
import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.chat import service
from app.chat.models import ChatTurn
from app.chat.sse import format_sse
from app.config import get_settings
from app.llm.fake import ScriptedChatModel
from app.llm.models import LlmProfile
from tests.llm.openai_mock import error_reply, mock_openai

FAKE_Q1 = "How many unique people attended Bagel Breakfast, October?"
FAKE_Q2 = "Of those, how many are international?"
EMAILS_Q = "List emails of everyone who attended any Bagel Breakfast"


@pytest.fixture
def fake_llm(monkeypatch):
    monkeypatch.setenv("HUB_ENABLE_FAKE_LLM", "true")
    get_settings.cache_clear()
    yield
    monkeypatch.undo()
    get_settings.cache_clear()


@pytest.fixture
def fake_profile(client, fake_llm, loaded_samples):
    res = client.post(
        "/api/llm-profiles", json={"name": "Demo", "preset": "fake", "model": "scripted"}
    )
    assert res.status_code == 201, res.text
    return res.json()["id"]


def sse_events(resp) -> list[tuple[str, dict]]:
    events = []
    for block in resp.read().decode().split("\n\n"):
        if not block.strip():
            continue
        lines = dict(line.split(": ", 1) for line in block.split("\n"))
        events.append((lines["event"], json.loads(lines["data"])))
    return events


def ask(client, cid, question, **body):
    with client.stream(
        "POST", f"/api/conversations/{cid}/messages", json={"question": question, **body}
    ) as r:
        assert r.status_code == 200, r.read()
        return sse_events(r)


def new_conversation(client) -> int:
    return client.post("/api/conversations").json()["id"]


def turns(client, cid) -> list[dict]:
    return client.get(f"/api/conversations/{cid}").json()["turns"]


def test_format_sse():
    assert format_sse("step", {"a": 1}) == 'event: step\ndata: {"a": 1}\n\n'


def test_stream_happy_path_and_turn_logged(client, fake_profile):
    cid = new_conversation(client)
    with client.stream(
        "POST",
        f"/api/conversations/{cid}/messages",
        json={"question": FAKE_Q1, "llm_profile_id": fake_profile},
    ) as r:
        assert r.headers["content-type"].startswith("text/event-stream")
        assert r.headers["x-accel-buffering"] == "no" and r.headers["cache-control"] == "no-cache"
        ev = sse_events(r)
    assert ev[0] == ("step", {"name": "generating", "attempt": 1})
    assert ev[-1][0] == "done"
    assert [e for e, _ in ev].count("error") == 0
    turn = turns(client, cid)[0]
    assert ev[-1][1] == {"turn_id": turn["id"]}
    assert turn["kind"] == "answer" and turn["sql"] and turn["model"] == "scripted"
    assert turn["seq"] == 1 and turn["llm_profile_id"] == fake_profile
    detail = client.get(f"/api/conversations/{cid}").json()
    assert detail["title"] == FAKE_Q1


def test_conversation_title_is_null_then_first_question_cut_to_80(client, fake_profile):
    created = client.post("/api/conversations")
    assert created.status_code == 201 and created.json()["title"] is None
    cid = created.json()["id"]
    ask(client, cid, "How many people attended? " + "x" * 200, llm_profile_id=fake_profile)
    ask(client, cid, FAKE_Q1, llm_profile_id=fake_profile)
    detail = client.get(f"/api/conversations/{cid}").json()
    assert len(detail["title"]) == 80 and detail["title"].startswith("How many people")
    assert [t["seq"] for t in detail["turns"]] == [1, 2]
    listed = client.get("/api/conversations").json()
    assert [c["id"] for c in listed] == [cid]


def test_follow_up_uses_history(client, fake_profile, monkeypatch):
    model = ScriptedChatModel.default()
    monkeypatch.setattr(service, "build_chat_model", lambda *a, **k: model)
    cid = new_conversation(client)
    first = ask(client, cid, FAKE_Q1, llm_profile_id=fake_profile)
    sql1 = next(d["sql"] for e, d in first if e == "sql")
    model.calls.clear()
    ask(client, cid, FAKE_Q2, llm_profile_id=fake_profile)
    prompt = "\n".join(str(m.content) for m in model.calls[0])
    assert sql1 in prompt


@pytest.mark.parametrize(
    "profile_share,request_share,source",
    [(True, True, "llm"), (True, False, "local"), (False, True, "local"), (False, False, "local")],
)
def test_sharing_needs_profile_and_request(
    client, fake_profile, profile_share, request_share, source
):
    client.patch(
        f"/api/llm-profiles/{fake_profile}", json={"share_results_with_llm": profile_share}
    )
    cid = new_conversation(client)
    ev = ask(client, cid, FAKE_Q1, llm_profile_id=fake_profile, share_results=request_share)
    analysis = next(d for e, d in ev if e == "analysis")
    assert analysis["source"] == source and analysis["results_shared"] == (source == "llm")
    assert turns(client, cid)[0]["results_shared"] == (source == "llm")


def test_clarification_turn(client, fake_profile):
    cid = new_conversation(client)
    ev = ask(client, cid, "How many people came this year?", llm_profile_id=fake_profile)
    assert [e for e, _ in ev][-2:] == ["clarification", "done"]
    turn = turns(client, cid)[0]
    assert turn["kind"] == "clarification" and turn["clarification_question"]


def test_rerun_and_export_xlsx_and_csv(client, fake_profile):
    cid = new_conversation(client)
    ask(client, cid, EMAILS_Q, llm_profile_id=fake_profile)
    turn = turns(client, cid)[0]
    rerun = client.get(f"/api/chat/turns/{turn['id']}/rerun")
    assert rerun.status_code == 200
    body = rerun.json()
    assert body["columns"] == ["email"] and body["row_count"] == turn["row_count"] > 0
    xlsx = client.get(f"/api/chat/turns/{turn['id']}/export?format=xlsx")
    assert (
        xlsx.status_code == 200 and f"hub-turn-{turn['id']}-" in xlsx.headers["content-disposition"]
    )
    sheet = load_workbook(io.BytesIO(xlsx.content)).active
    assert sheet.max_row - 1 == body["row_count"]
    csv = client.get(f"/api/chat/turns/{turn['id']}/export?format=csv")
    assert csv.status_code == 200 and csv.text.lstrip("\ufeff").splitlines()[0] == "email"


def test_export_formula_guard(client, fake_profile, db_session, user):
    cid = new_conversation(client)
    ask(client, cid, FAKE_Q1, llm_profile_id=fake_profile)
    turn = db_session.scalar(select(ChatTurn))
    turn.sql = "SELECT '=1+1' AS v FROM attendance_view LIMIT 1"
    db_session.commit()
    csv = client.get(f"/api/chat/turns/{turn.id}/export?format=csv")
    assert "'=1+1" in csv.text


def test_rerun_and_export_errors(client, fake_profile):
    assert client.get("/api/chat/turns/9999/rerun").status_code == 404
    cid = new_conversation(client)
    ask(client, cid, "How many people came this year?", llm_profile_id=fake_profile)
    tid = turns(client, cid)[0]["id"]
    for path in ("rerun", "export?format=csv"):
        res = client.get(f"/api/chat/turns/{tid}/{path}")
        assert res.status_code == 409 and res.json()["error"]["code"] == "turn_has_no_result"
    assert client.get(f"/api/chat/turns/{tid}/export?format=pdf").status_code == 422


def test_message_to_missing_conversation_is_404(client, fake_profile):
    res = client.post("/api/conversations/999/messages", json={"question": "hi"})
    assert res.status_code == 404 and res.json()["error"]["code"] == "conversation_not_found"


def error_turn(client, cid, **body):
    ev = ask(client, cid, "How many people attended?", **body)
    assert [e for e, _ in ev].count("error") == 1 and ev[-1][0] == "done"
    return next(d for e, d in ev if e == "error"), turns(client, cid)[-1]


def test_profile_deleted_gives_error_event(client):
    err, turn = error_turn(client, new_conversation(client), llm_profile_id=999)
    assert err["code"] == "profile_not_found" and err["message"]
    assert turn["kind"] == "error" and turn["error_code"] == "profile_not_found"
    assert turn["llm_profile_id"] is None


def test_missing_key_gives_error_event(client):
    created = client.post(
        "/api/llm-profiles", json={"name": "NoKey", "preset": "openai", "model": "gpt-4.1-mini"}
    ).json()
    err, turn = error_turn(client, new_conversation(client), llm_profile_id=created["id"])
    assert err["code"] == "profile_missing_key"
    assert turn["kind"] == "error" and turn["llm_profile_id"] == created["id"]


@pytest.mark.parametrize(
    "status,code",
    [(401, "llm_auth_failed"), (429, "llm_rate_limited"), (404, "llm_model_not_found")],
)
def test_provider_errors_become_error_events(client, loaded_samples, monkeypatch, status, code):
    from app.llm.factory import build_chat_model as real

    def patched(profile, key, *, settings, http_async_client=None):
        client_ = mock_openai(lambda request: error_reply(status, "nope sk-secret"))
        return real(profile, key, settings=settings, http_async_client=client_)

    monkeypatch.setattr(service, "build_chat_model", patched)
    created = client.post(
        "/api/llm-profiles",
        json={"name": "OA", "preset": "openai", "model": "m", "api_key": "sk-secret-1234567890"},
    ).json()
    err, turn = error_turn(client, new_conversation(client), llm_profile_id=created["id"])
    assert err["code"] == code and "sk-secret" not in json.dumps(err)
    assert turn["kind"] == "error" and turn["error_code"] == code


def test_json_mode_detection_is_persisted(client, loaded_samples, db_session, monkeypatch):
    from app.llm.factory import build_chat_model as real
    from tests.llm.openai_mock import chat_reply, request_json

    sql = json.dumps(
        {"kind": "sql", "sql": "SELECT COUNT(*) AS n FROM attendance_view", "interpretation": "x"}
    )

    def handler(request: httpx.Request) -> httpx.Response:
        if "response_format" in request_json(request):
            return error_reply(400, "Unsupported parameter: response_format")
        return chat_reply(sql)

    monkeypatch.setattr(
        service,
        "build_chat_model",
        lambda p, k, *, settings: real(
            p, k, settings=settings, http_async_client=mock_openai(handler)
        ),
    )
    created = client.post(
        "/api/llm-profiles",
        json={"name": "OA", "preset": "openai", "model": "m", "api_key": "sk-secret-1234567890"},
    ).json()
    ask(client, new_conversation(client), "How many rows?", llm_profile_id=created["id"])
    db_session.expire_all()
    assert db_session.get(LlmProfile, created["id"]).json_mode is False


def test_turn_timeout_gives_error_event(client, fake_profile, monkeypatch):
    monkeypatch.setenv("TURN_TIMEOUT_S", "0.05")
    get_settings.cache_clear()
    model = ScriptedChatModel.default()

    async def slow(*a, **k):
        import asyncio

        await asyncio.sleep(5)

    monkeypatch.setattr(model.__class__, "_agenerate", slow)
    monkeypatch.setattr(service, "build_chat_model", lambda *a, **k: model)
    cid = new_conversation(client)
    err, turn = error_turn(client, cid, llm_profile_id=fake_profile)
    assert err["code"] == "turn_timeout" and turn["error_code"] == "turn_timeout"
    get_settings.cache_clear()


def test_unexpected_exception_becomes_internal_error(client, fake_profile, monkeypatch):
    def boom(engine):
        raise RuntimeError("secret row value jane@uci.edu")

    monkeypatch.setattr("app.chat.graph.nodes.load_controlled_values", boom)
    err, turn = error_turn(client, new_conversation(client), llm_profile_id=fake_profile)
    assert err["code"] == "internal_error" and "jane" not in json.dumps(err)
    assert turn["kind"] == "error" and turn["error_code"] == "internal_error"


def test_question_too_long_is_422(client, fake_profile):
    cid = new_conversation(client)
    res = client.post(
        f"/api/conversations/{cid}/messages",
        json={"question": "x" * 2001, "llm_profile_id": fake_profile},
    )
    assert res.status_code == 422 and res.json()["error"]["code"] == "validation_error"
    res = client.post(f"/api/conversations/{cid}/messages", json={"question": ""})
    assert res.status_code == 422


def slow_model(monkeypatch, *, only_analysis: bool):
    import asyncio

    from app.chat.graph.prompts import ANALYZE_SYSTEM

    model = ScriptedChatModel.default()
    real = ScriptedChatModel._agenerate

    async def slow(self, messages, *a, **k):
        if not only_analysis or any(m.content == ANALYZE_SYSTEM for m in messages):
            await asyncio.sleep(10)
        return await real(self, messages, *a, **k)

    monkeypatch.setattr(ScriptedChatModel, "_agenerate", slow)
    monkeypatch.setattr(service, "build_chat_model", lambda *a, **k: model)


def test_timeout_during_analysis_keeps_partial_turn(client, fake_profile, monkeypatch):
    monkeypatch.setenv("TURN_TIMEOUT_S", "2")
    get_settings.cache_clear()
    slow_model(monkeypatch, only_analysis=True)
    cid = new_conversation(client)
    ev = ask(client, cid, FAKE_Q1, llm_profile_id=fake_profile, share_results=True)
    assert [e for e, _ in ev].count("error") == 1 and ev[-1][0] == "done"
    assert next(d for e, d in ev if e == "error")["code"] == "turn_timeout"
    turn = turns(client, cid)[-1]
    assert turn["error_code"] == "turn_timeout" and ev[-1][1] == {"turn_id": turn["id"]}
    assert turn["sql"] and turn["row_count"] == 1 and turn["results_shared"] is True
    assert turns(client, cid) == [turn]
    get_settings.cache_clear()


def test_timeout_while_recording_writes_one_turn_and_done_has_id(client, fake_profile, monkeypatch):
    import time

    monkeypatch.setenv("TURN_TIMEOUT_S", "2")
    get_settings.cache_clear()
    real = service._write_turn

    def slow_write(*a, **k):
        time.sleep(3)
        return real(*a, **k)

    monkeypatch.setattr(service, "_write_turn", slow_write)
    cid = new_conversation(client)
    ev = ask(client, cid, FAKE_Q1, llm_profile_id=fake_profile)
    rows = turns(client, cid)
    assert len(rows) == 1 and ev[-1] == ("done", {"turn_id": rows[0]["id"]})
    get_settings.cache_clear()


def test_other_timeout_inside_graph_is_internal_error(client, fake_profile, monkeypatch):
    def boom(engine):
        raise TimeoutError

    monkeypatch.setattr("app.chat.graph.nodes.load_controlled_values", boom)
    err, _ = error_turn(client, new_conversation(client), llm_profile_id=fake_profile)
    assert err["code"] == "internal_error"


def test_prepared_repr_hides_key():
    prepared = service._Prepared(
        profile_id=1, preset="openai", model_name="m", share=True, json_mode=None,
        api_key="sk-secret-1234567890", base_url="https://x", history=[],
    )  # fmt: skip
    assert "sk-secret" not in repr(prepared)


async def test_orphaned_failed_write_logs_no_values(
    client, fake_profile, user, monkeypatch, caplog
):
    import asyncio
    import gc
    import threading
    import time

    from app.chat.schemas import MessageIn

    started = threading.Event()

    def failing_write(*a, **k):
        started.set()
        time.sleep(0.3)
        raise RuntimeError("[parameters: SECRET-QUESTION SECRET-SQL SECRET-ANALYSIS]")

    monkeypatch.setattr(service, "_write_turn", failing_write)
    cid = new_conversation(client)
    gen = service.run_turn(
        cid, MessageIn(question=FAKE_Q1, llm_profile_id=fake_profile, share_results=True), user
    )

    async def consume():
        async for _ in gen:
            pass

    task = asyncio.create_task(consume())
    while not started.is_set():
        await asyncio.sleep(0.01)
    task.cancel()  # a client disconnect
    with pytest.raises(asyncio.CancelledError):
        await task
    await asyncio.sleep(0.6)
    gc.collect()
    await asyncio.sleep(0)
    assert "never retrieved" not in caplog.text and "SECRET" not in caplog.text
    assert "Recording the turn failed (RuntimeError)" in caplog.text


def _profile_with_unreadable_key(client, db_session):
    from cryptography.fernet import Fernet

    from app.llm.models import LlmProfile

    created = client.post(
        "/api/llm-profiles",
        json={"name": "Rotated", "preset": "openai", "model": "m", "api_key": "sk-test-1234567890"},
    ).json()
    row = db_session.get(LlmProfile, created["id"])
    row.api_key_ciphertext = Fernet(Fernet.generate_key()).encrypt(b"sk-old")  # another secret key
    db_session.commit()
    return created["id"]


def test_unreadable_saved_key_gives_error_event(client, db_session):
    pid = _profile_with_unreadable_key(client, db_session)
    err, turn = error_turn(client, new_conversation(client), llm_profile_id=pid)
    assert err["code"] == "profile_missing_key" and "secret key changed" in err["message"]
    assert turn["error_code"] == "profile_missing_key"


def test_unreadable_saved_key_gives_409_on_test_route(client, db_session):
    pid = _profile_with_unreadable_key(client, db_session)
    res = client.post(f"/api/llm-profiles/{pid}/test")
    assert res.status_code == 409 and res.json()["error"]["code"] == "profile_missing_key"


def _stream_turn(cid, user, fake_profile, **body):
    from app.chat.schemas import MessageIn

    return service.run_turn(
        cid, MessageIn(question=FAKE_Q1, llm_profile_id=fake_profile, **body), user
    )


def _only_turn(db_session) -> ChatTurn:
    db_session.expire_all()
    return db_session.scalars(select(ChatTurn)).one()


async def test_disconnect_mid_stream_records_partial_turn(
    client, fake_profile, user, db_session, monkeypatch
):
    import asyncio

    slow_model(monkeypatch, only_analysis=True)
    cid = new_conversation(client)
    gen = _stream_turn(cid, user, fake_profile, share_results=True)
    seen: list[str] = []

    async def consume():
        async for chunk in gen:
            seen.append(chunk)

    task = asyncio.create_task(consume())
    while not any('"analyzing"' in c for c in seen):
        await asyncio.sleep(0.01)
    task.cancel()  # the client went away
    with pytest.raises(asyncio.CancelledError):
        await task
    turn = _only_turn(db_session)
    assert turn.kind == "error" and turn.error_code == "client_disconnected"
    assert turn.sql and turn.row_count == 1 and turn.results_shared is True
    assert not any("event: error" in c or "event: done" in c for c in seen)


async def test_disconnect_at_a_yield_records_one_turn(client, fake_profile, user, db_session):
    cid = new_conversation(client)
    gen = _stream_turn(cid, user, fake_profile)
    await anext(gen)  # the first event; the generator is now suspended at its yield
    await gen.aclose()
    turn = _only_turn(db_session)
    assert turn.error_code == "client_disconnected"


async def test_disconnect_after_done_adds_no_second_turn(client, fake_profile, user, db_session):
    cid = new_conversation(client)
    gen = _stream_turn(cid, user, fake_profile)
    chunks = []
    async for chunk in gen:
        chunks.append(chunk)
        if chunk.startswith("event: done"):
            break
    await gen.aclose()
    turn = _only_turn(db_session)
    assert turn.kind == "answer" and turn.error_code is None


def test_export_reports_truncation_headers(client, fake_profile, monkeypatch):
    cid = new_conversation(client)
    ask(client, cid, EMAILS_Q, llm_profile_id=fake_profile)
    turn = turns(client, cid)[0]
    url = f"/api/chat/turns/{turn['id']}/export"
    full = client.get(f"{url}?format=csv")
    assert full.headers["x-hub-truncated"] == "false"
    assert full.headers["x-hub-row-cap"] == "100000"
    monkeypatch.setenv("EXPORT_ROW_CAP", "3")
    get_settings.cache_clear()
    try:
        cut = client.get(f"{url}?format=csv")
        assert cut.headers["x-hub-truncated"] == "true" and cut.headers["x-hub-row-cap"] == "3"
        assert len(cut.text.lstrip("﻿").splitlines()) == 1 + 3
        monkeypatch.setenv("EXPORT_ROW_CAP", str(turn["row_count"]))
        get_settings.cache_clear()
        exact = client.get(f"{url}?format=xlsx")  # exactly at the cap is not truncated
        assert exact.headers["x-hub-truncated"] == "false"
    finally:
        monkeypatch.undo()
        get_settings.cache_clear()


def test_whitespace_only_question_is_422_and_question_is_stripped(client, fake_profile):
    cid = new_conversation(client)
    url = f"/api/conversations/{cid}/messages"
    for blank in ("   ", "\n\t "):
        res = client.post(url, json={"question": blank})
        assert res.status_code == 422 and res.json()["error"]["code"] == "validation_error"
    # the limit applies after stripping: 2,000 characters plus padding is fine
    padded = f"  {'x' * 2000}  "
    res = client.post(url, json={"question": padded, "llm_profile_id": fake_profile})
    assert res.status_code == 200
    res = client.post(url, json={"question": "x" * 2001 + "  "})
    assert res.status_code == 422
    assert turns(client, cid)[0]["question"] == "x" * 2000
