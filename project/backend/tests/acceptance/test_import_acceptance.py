from collections import Counter
from datetime import date, datetime
from io import BytesIO

import pytest
from openpyxl import load_workbook
from sqlalchemy import text

from app.dataset.service import gpsrc_layout_rows
from app.export.writers import to_xlsx
from app.imports.service import preview_upload
from app.samples import SAMPLE_DIR, load_sample_exports

pytestmark = pytest.mark.acceptance

COMPILED = SAMPLE_DIR.parent / "GSPRC Compiled Dataset.xlsm"

COMPARE = [
    "Student Worker", "Event Term", "Event Name ", "Virtual/ In-Person/ Hybrid ",
    "Sponsor/Host/Parnters",
    "Event type (Professional Development, Wellness/Mental Health, Fellowships)",
    "On Campus Groups?", "First Name of Attendee", "Last Name of Attendee", "Email",
    "Affiliation", "Student ID or Employee ID", "NetID", "Domestic or International", "School",
    "Department ",
]  # fmt: skip


def norm(v) -> str:
    if v is None or v == "":
        return ""
    if v is True:
        return "Yes"
    if v is False:
        return "No"
    if isinstance(v, (int, float)):
        return str(int(v)) if float(v).is_integer() else str(v)
    if isinstance(v, (date, datetime)):
        return v.isoformat()
    return str(v).strip()


def compiled_rows(path, columns):
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = wb["Master Attendance"].iter_rows(values_only=True)
        header = list(next(rows))
        idx = [header.index(c) for c in columns]
        out = []
        for row in rows:
            row = list(row) + [None] * (len(header) - len(row))
            if all(v is None or v == "" for v in row):
                continue
            out.append(tuple(norm(row[i]) for i in idx))
        return out
    finally:
        wb.close()


def _diff_message(ours: Counter, theirs: Counter) -> str:
    only_ours = list((ours - theirs).elements())[:5]
    only_theirs = list((theirs - ours).elements())[:5]
    return (
        f"ours={sum(ours.values())} theirs={sum(theirs.values())}\n"
        f"only in import ({sum((ours - theirs).values())}): {only_ours}\n"
        f"only in compiled sheet ({sum((theirs - ours).values())}): {only_theirs}"
    )


def test_a1_import_reproduces_compiled_sheet(db_session, user):
    load_sample_exports(db_session, user)
    cols, rows = gpsrc_layout_rows(db_session)
    ours = Counter(tuple(norm(r[cols.index(c)]) for c in COMPARE) for r in rows)
    theirs = Counter(compiled_rows(COMPILED, COMPARE))
    assert sum(ours.values()) == 293 and ours == theirs, _diff_message(ours, theirs)
    assert db_session.scalar(text("SELECT count(DISTINCT netid) FROM hub.attendance_view")) == 79


def test_a1_identity_edge_rows(db_session, user):
    load_sample_exports(db_session, user)
    nov = db_session.execute(
        text(
            "SELECT a.match_method, a.person_id FROM hub.attendance a "
            "JOIN hub.event e ON e.id = a.event_id "
            "WHERE e.name = 'Bagel Breakfast, November' AND a.email = 'cdiaz@uci.edu'"
        )
    ).one()
    assert nov.match_method == "email"
    october = db_session.scalar(
        text(
            "SELECT a.person_id FROM hub.attendance a JOIN hub.event e ON e.id = a.event_id "
            "WHERE e.name = 'Bagel Breakfast, October' AND a.netid = 'cdiaz'"
        )
    )
    assert october is not None and nov.person_id == october

    rows = db_session.execute(
        text(
            "SELECT a.id, a.person_id FROM hub.attendance a JOIN hub.event e ON e.id = a.event_id "
            "WHERE e.name = 'LinkedIn Info Session' "
            "AND coalesce(a.email, '') = '' AND coalesce(a.netid, '') = ''"
        )
    ).all()
    assert len(rows) == 1 and rows[0].person_id is None
    flags = db_session.scalar(
        text(
            "SELECT count(*) FROM hub.data_quality_flag "
            "WHERE attendance_id = :id AND code = 'missing_identity'"
        ),
        {"id": rows[0].id},
    )
    assert flags == 1


def test_a2_reupload_adds_nothing(db_session, user):
    load_sample_exports(db_session, user)
    for f in sorted(SAMPLE_DIR.iterdir()):
        assert preview_upload(db_session, f.name, f.read_bytes(), user).status == "duplicate"
    assert db_session.scalar(text("SELECT count(*) FROM hub.attendance")) == 293


def test_a9_formula_cell_exported_as_text():
    data = to_xlsx(["Name"], [["=1+1"], ["+SUM(A1)"], ["@cmd"]])
    ws = load_workbook(BytesIO(data)).active
    for row in range(2, 5):
        cell = ws.cell(row=row, column=1)
        assert cell.data_type == "s"
