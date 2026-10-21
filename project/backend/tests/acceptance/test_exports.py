import io
import json

import pytest
from openpyxl import load_workbook

from app.llm.fake import SCRIPT_PATH
from tests.acceptance.test_privacy import fake_llm, make_profile  # noqa: F401

pytestmark = pytest.mark.acceptance

# Every fake-script entry that answers with SQL; its match text doubles as the question.
ANSWER_QUESTIONS = [
    e["match"]
    for e in json.loads(SCRIPT_PATH.read_text(encoding="utf-8"))
    if json.loads(e["generate"][0])["kind"] == "sql"
]


def rows_in(fmt: str, resp) -> int:
    if fmt == "xlsx":
        return load_workbook(io.BytesIO(resp.content)).active.max_row - 1
    return len(resp.content.decode("utf-8-sig").splitlines()) - 1


def assert_attachment(fmt: str, resp) -> None:
    assert resp.status_code == 200
    assert resp.headers["content-disposition"].startswith("attachment;")
    assert resp.headers["content-disposition"].rstrip('"').endswith(f".{fmt}")


@pytest.mark.parametrize("fmt", ["xlsx", "csv"])
def test_a5_every_result_exports(client, loaded_samples, fake_llm, fmt):  # noqa: F811
    assert len(ANSWER_QUESTIONS) >= 3
    profile_id = make_profile(client, False)
    cid = client.post("/api/conversations").json()["id"]
    for question in ANSWER_QUESTIONS:
        with client.stream(
            "POST",
            f"/api/conversations/{cid}/messages",
            json={"question": question, "llm_profile_id": profile_id},
        ) as r:
            assert "event: error" not in r.read().decode()
    turns = client.get(f"/api/conversations/{cid}").json()["turns"]
    assert [t["kind"] for t in turns] == ["answer"] * len(ANSWER_QUESTIONS)
    for turn in turns:
        rerun = client.get(f"/api/chat/turns/{turn['id']}/rerun").json()
        export = client.get(f"/api/chat/turns/{turn['id']}/export?format={fmt}")
        assert_attachment(fmt, export)
        assert rows_in(fmt, export) == rerun["row_count"] > 0

    dataset = client.get(f"/api/dataset/export?format={fmt}")
    assert_attachment(fmt, dataset)
    assert rows_in(fmt, dataset) == 293
