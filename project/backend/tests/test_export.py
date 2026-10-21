from datetime import date, datetime
from io import BytesIO

import openpyxl
import pytest

from app.dataset.service import GPSRC_HEADERS
from app.export.writers import export_filename, export_response, guard_cell, to_csv, to_xlsx


def test_formula_value_is_text_in_xlsx():
    wb = openpyxl.load_workbook(BytesIO(to_xlsx(["Department"], [['=HYPERLINK("x")']])))
    c = wb.active["A2"]
    assert c.data_type == "s" and c.value == '=HYPERLINK("x")'


def test_formula_value_prefixed_in_csv():
    assert to_csv(["d"], [["+1"]]).decode("utf-8-sig").splitlines()[1] == "'+1"


def test_plain_values_untouched():
    assert guard_cell("Biology") == "Biology" and guard_cell(5) == 5
    assert guard_cell(None) is None and guard_cell(date(2026, 1, 2)) == date(2026, 1, 2)


@pytest.mark.parametrize("trigger", ["=", "+", "-", "@", "\t", "\r"])
def test_guard_cell_each_trigger(trigger):
    assert guard_cell(f"{trigger}x") == f"'{trigger}x"
    assert guard_cell(f"a{trigger}") == f"a{trigger}"


def test_xlsx_none_and_dates():
    wb = openpyxl.load_workbook(BytesIO(to_xlsx(["a", "b"], [[None, date(2026, 10, 22)]])))
    ws = wb.active
    assert ws["A2"].value is None
    assert ws["B2"].value == datetime(2026, 10, 22)


def test_csv_none_dates_and_bom():
    raw = to_csv(["a", "b"], [[None, date(2026, 10, 22)]])
    assert raw.startswith(b"\xef\xbb\xbf")
    assert raw.decode("utf-8-sig").splitlines()[1] == ",2026-10-22"


def test_export_filename():
    assert (
        export_filename("dataset", "xlsx", datetime(2026, 10, 22, 9, 5))
        == "hub-dataset-20261022-0905.xlsx"
    )


def test_export_response_headers():
    r = export_response(["a"], [["x"]], "csv", "dataset")
    assert r.media_type == "text/csv; charset=utf-8"
    assert r.headers["content-disposition"].startswith('attachment; filename="hub-dataset-')
    assert r.headers["content-disposition"].endswith('.csv"')
    x = export_response(["a"], [["x"]], "xlsx", "dataset")
    assert x.media_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def test_full_export_headers(client, committed_october):
    r = client.get("/api/dataset/export?format=xlsx")
    assert r.status_code == 200
    ws = openpyxl.load_workbook(BytesIO(r.content)).active
    assert [c.value for c in ws[1]] == list(GPSRC_HEADERS) and ws.max_row == 39


def test_full_export_default_is_xlsx(client, committed_october):
    r = client.get("/api/dataset/export")
    assert r.headers["content-disposition"].endswith('.xlsx"')


def test_full_export_csv(client, committed_october):
    r = client.get("/api/dataset/export?format=csv")
    assert r.status_code == 200
    lines = r.content.decode("utf-8-sig").splitlines()
    assert len(lines) == 39
    import csv

    assert next(csv.reader([lines[0]])) == list(GPSRC_HEADERS)


def test_export_rejects_bad_format(client):
    r = client.get("/api/dataset/export?format=pdf")
    assert r.status_code == 422 and r.json()["error"]["code"] == "validation_error"


def test_xlsx_drops_control_characters_instead_of_failing():
    data = to_xlsx(["Department\x0b"], [["Bio\x0bChem", "=x\x01", 5, None]])
    ws = openpyxl.load_workbook(BytesIO(data)).active
    assert ws["A1"].value == "Department"
    assert ws["A2"].value == "BioChem"
    assert ws["B2"].value == "=x" and ws["B2"].data_type == "s"


def test_csv_keeps_values_as_given():
    # str.splitlines() would also split on \x0b, so split on the csv module's terminator.
    assert to_csv(["d"], [["Bio\x0bChem"]]).decode("utf-8-sig").split("\r\n")[1] == "Bio\x0bChem"


def test_export_still_works_after_importing_a_control_character(client):
    from tests.imports.helpers import CG_HEADER, ROW_AAHMED, csv_with_rows

    row = list(ROW_AAHMED)
    row[CG_HEADER.index("Department")] = "Bio\x0bChem"
    files = [("files", ("vt.csv", csv_with_rows(CG_HEADER, [row]), "text/csv"))]
    import_id = client.post("/api/imports", files=files).json()["files"][0]["id"]
    details = {"name": "VT", "term": "Fall 2026", "modality": "In Person", "event_type": "Other"}
    assert client.post(f"/api/imports/{import_id}/commit", json=details).status_code == 200

    r = client.get("/api/dataset/export?format=xlsx")
    assert r.status_code == 200
    ws = openpyxl.load_workbook(BytesIO(r.content)).active
    assert ws.cell(row=2, column=list(GPSRC_HEADERS).index("Department ") + 1).value == "BioChem"
