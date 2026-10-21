import csv
import io
from datetime import date, datetime
from pathlib import PurePath

from openpyxl import load_workbook

from app.errors import HubError

Grid = list[list[str | None]]


def _unreadable(message: str) -> HubError:
    return HubError("unreadable_file", message, 422)


def _cell(v) -> str | None:
    if v is None:
        return None
    if isinstance(v, bool):
        return str(v)
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() else str(v)
    if isinstance(v, int):
        return str(v)
    if isinstance(v, datetime | date):
        return v.isoformat()
    return str(v).strip() or None


def _shape(rows: list[list]) -> Grid:
    cleaned = [[_cell(c) for c in row] for row in rows]
    cleaned = [row for row in cleaned if any(c is not None for c in row)]
    if not cleaned:
        return []
    header = cleaned[0]
    width = max(i for i, c in enumerate(header) if c is not None) + 1
    return [(row + [None] * width)[:width] for row in cleaned]


def _read_csv(content: bytes) -> list[list]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = content.decode("cp1252", errors="replace")
    return list(csv.reader(io.StringIO(text, newline="")))


def _read_xlsx(content: bytes) -> list[list]:
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        try:
            ws = wb.worksheets[0]
            return [list(row) for row in ws.iter_rows(values_only=True)]
        finally:
            wb.close()
    except HubError:
        raise
    except Exception as exc:
        raise _unreadable("The file could not be read as a spreadsheet.") from exc


def read_grid(file_name: str, content: bytes) -> Grid:
    ext = PurePath(file_name).suffix.lower()
    if ext == ".csv":
        rows = _read_csv(content)
    elif ext in (".xlsx", ".xlsm"):
        rows = _read_xlsx(content)
    else:
        raise _unreadable("Unsupported file type; upload a .csv, .xlsx, or .xlsm file.")
    return _shape(rows)
