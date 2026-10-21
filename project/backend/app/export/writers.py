import csv
import io
from collections.abc import Iterable, Sequence
from datetime import date, datetime
from typing import Literal

from fastapi import Response
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

XLSX_MEDIA = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
CSV_MEDIA = "text/csv; charset=utf-8"
_TRIGGERS = ("=", "+", "-", "@", "\t", "\r")


def _needs_guard(value: object) -> bool:
    return isinstance(value, str) and value.startswith(_TRIGGERS)


def guard_cell(value: object) -> object:
    """Neutralise spreadsheet formula injection: mark trigger-leading strings as text."""
    return f"'{value}" if _needs_guard(value) else value


def _xlsx_value(value: object) -> object:
    """Drop control characters openpyxl refuses to write (import already strips them)."""
    return ILLEGAL_CHARACTERS_RE.sub("", value) if isinstance(value, str) else value


def to_xlsx(columns: list[str], rows: Iterable[Sequence[object]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append([_xlsx_value(c) for c in columns])
    for row in rows:
        ws.append([_xlsx_value(v) for v in row])
        for cell in ws[ws.max_row]:
            if _needs_guard(cell.value):
                cell.data_type = "s"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _csv_value(value: object) -> object:
    if value is None:
        return ""
    if isinstance(value, date | datetime):
        return value.isoformat()
    return guard_cell(value)


def to_csv(columns: list[str], rows: Iterable[Sequence[object]]) -> bytes:
    buf = io.StringIO(newline="")
    writer = csv.writer(buf)
    writer.writerow(list(columns))
    for row in rows:
        writer.writerow([_csv_value(v) for v in row])
    return buf.getvalue().encode("utf-8-sig")


def export_filename(stem: str, ext: str, now: datetime) -> str:
    return f"hub-{stem}-{now:%Y%m%d-%H%M}.{ext}"


def export_response(
    columns: list[str],
    rows: Iterable[Sequence[object]],
    fmt: Literal["xlsx", "csv"],
    stem: str,
) -> Response:
    if fmt == "xlsx":
        body, media = to_xlsx(columns, rows), XLSX_MEDIA
    else:
        body, media = to_csv(columns, rows), CSV_MEDIA
    name = export_filename(stem, fmt, datetime.now().astimezone())
    return Response(
        content=body,
        media_type=media,
        headers={"Content-Disposition": f'attachment; filename="{name}"'},
    )
